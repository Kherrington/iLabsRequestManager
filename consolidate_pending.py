"""
Merge the temporary time-stamped files the app writes when a shared xlsx is
locked back into the real workbooks, without duplicating any rows.

The app saves rows it could not write to a locked file as e.g.
    Training Schedule_CALM_CURRENT.pending_20261021_101500.xlsx
next to the real file.  This script appends every row of those files that is
not already present, then moves the processed files into a "pending_merged"
sub-folder (nothing is deleted).  Files whose target is still locked are left
in place; run the script again later.

Covers the CALM schedule, the CVRI schedule and the Microscope Intro Course Log.

Usage:
    python consolidate_pending.py             # merge everything
    python consolidate_pending.py --dry-run   # only report what would happen
"""

import argparse
import datetime as _dt
import shutil
import sys
from pathlib import Path

import openpyxl

import prefs
from xlsx_export import HEADER_MAP, _norm, pending_files

# (label, prefs key of the xlsx, prefs key of its sheet name)
TARGETS = [
    ("CALM schedule",     "calm_xlsx",  "calm_sheet"),
    ("CVRI schedule",     "cvri_xlsx",  "cvri_sheet"),
    ("Intro Course Log",  "intro_xlsx", "intro_sheet"),
]


def _cell_text(v) -> str:
    """Normalise a cell value for comparison (case-insensitive, trimmed;
    dates compare by ISO date whatever way they were typed/stored)."""
    if v is None:
        return ""
    if isinstance(v, _dt.datetime):
        return (v.date().isoformat() if v.time() == _dt.time(0)
                else v.isoformat(timespec="minutes"))
    if isinstance(v, _dt.date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    text = " ".join(str(v).split())
    for fmt in ("%m/%d/%Y", "%m/%d/%y", "%Y-%m-%d"):
        try:
            return _dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return text.lower()


def _headers(ws) -> list:
    raw = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    headers = [str(h).strip() if h is not None else "" for h in raw]
    while headers and not headers[-1]:
        headers.pop()
    return headers


def _sheet(wb, name: str):
    return wb[name] if name and name in wb.sheetnames else wb.active


def _merge_target(label: str, xlsx: Path, sheet: str, dry_run: bool) -> None:
    pend = pending_files(str(xlsx))
    if not pend:
        print(f"{label}: no temporary files.")
        return
    if not xlsx.exists():
        print(f"{label}: {xlsx.name} not found — skipped.")
        return

    wb = openpyxl.load_workbook(xlsx)
    ws = _sheet(wb, sheet)
    headers = _headers(ws)
    if not headers:
        print(f"{label}: {xlsx.name} has no header row — skipped.")
        return
    norm_headers = [_norm(h) for h in headers]
    n = len(headers)

    # Existing rows: full-row signatures, plus Request IDs where there is one
    sigs, req_ids = set(), set()
    req_col = next((i for i, h in enumerate(headers)
                    if h and HEADER_MAP.get(_norm(h)) == "request_id"), None)
    for r in range(2, ws.max_row + 1):
        vals = tuple(_cell_text(ws.cell(r, c + 1).value) for c in range(n))
        if any(vals):
            sigs.add(vals)
            if req_col is not None and vals[req_col]:
                req_ids.add(vals[req_col])

    added = dup = 0
    for pf in pend:
        pwb = openpyxl.load_workbook(pf, read_only=True, data_only=True)
        try:
            pws = pwb.active
            rows = list(pws.iter_rows(values_only=True))
        finally:
            pwb.close()
        if not rows:
            continue
        pidx = {_norm(str(h)): i for i, h in enumerate(rows[0]) if h is not None}
        for prow in rows[1:]:
            # Line the pending columns up with the real file's headers
            aligned = [prow[pidx[h]] if h in pidx and pidx[h] < len(prow) else None
                       for h in norm_headers]
            sig = tuple(_cell_text(v) for v in aligned)
            if not any(sig):
                continue
            rid = sig[req_col] if req_col is not None else ""
            if sig in sigs or (rid and rid in req_ids):
                dup += 1
                continue
            if not dry_run:
                ws.append(aligned)
            sigs.add(sig)
            if rid:
                req_ids.add(rid)
            added += 1
    print(f"{label}: {len(pend)} temporary file(s) — "
          f"{added} row(s) {'would be ' if dry_run else ''}added, "
          f"{dup} duplicate(s) skipped.")

    if dry_run:
        return
    try:
        if added:
            wb.save(xlsx)
    except PermissionError:
        print(f"  {xlsx.name} is locked (open in Excel?) — temporary files "
              "left in place; run again later.")
        return
    done = xlsx.parent / "pending_merged"
    done.mkdir(exist_ok=True)
    for pf in pend:
        shutil.move(str(pf), str(done / pf.name))
    print(f"  Moved {len(pend)} file(s) to {done.name}\\")


def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    ap.add_argument("--dry-run", action="store_true",
                    help="report what would be merged without changing anything")
    args = ap.parse_args()

    p = prefs.get_prefs()
    if not prefs.shared_dir(p).is_dir():
        print(f"Shared folder not reachable: {prefs.shared_dir(p)}")
        return 1
    for label, xkey, skey in TARGETS:
        try:
            _merge_target(label, Path(prefs.abs_path(p[xkey])),
                          str(p.get(skey, "") or "").strip(), args.dry_run)
        except Exception as exc:        # keep going with the other files
            print(f"{label}: failed — {exc}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
